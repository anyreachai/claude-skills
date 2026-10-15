from pptx import Presentation
from pptx.util import Emu
from status_deck.theme import Theme
from status_deck.slides import progress_bars

def _s():
    p = Presentation()
    p.slide_width = Emu(12192000)
    p.slide_height = Emu(6858000)
    return p.slides.add_slide(p.slide_layouts[6])

def test_progress_bars_render():
    s = _s()
    progress_bars.render(s, {"title": "Progress", "items": [
        {"label": "Order status", "build": 95, "tested": 90, "note": "holds at 90"}]}, Theme.anyreach())
    texts = [sh.text_frame.text for sh in s.shapes if sh.has_text_frame]
    assert "Order status" in texts and "95%" in texts and "90%" in texts


def test_progress_bar_fills_use_series_colours():
    s = _s()
    theme = Theme.anyreach()
    progress_bars.render(s, {"title": "Progress", "items": [
        {"label": "Order status", "build": 95, "tested": 90}]}, theme)
    fills = set()
    for sh in s.shapes:
        try:
            fills.add(sh.fill.fore_color.rgb)
        except Exception:
            pass
    # build bar is series 0 (indigo), tested bar is series 1 (purple)
    assert theme.series_color(0) in fills
    assert theme.series_color(1) in fills
