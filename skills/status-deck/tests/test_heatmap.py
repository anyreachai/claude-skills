from pptx import Presentation
from pptx.util import Emu
from status_deck.theme import Theme
from status_deck.slides import heatmap

def _s():
    p = Presentation()
    p.slide_width = Emu(12192000)
    p.slide_height = Emu(6858000)
    return p.slides.add_slide(p.slide_layouts[6])

def test_heatmap_cells():
    s = _s()
    heatmap.render(s, {"title": "Readiness", "row_labels": ["1.1", "1.2"], "col_labels": ["Build", "Test"],
                     "cells": [["green", "amber"], ["amber", "red"]]}, Theme.anyreach())
    fills = set()
    for sh in s.shapes:
        try:
            if sh.shape_type is not None and "TEXT" not in str(sh.shape_type):
                fills.add(str(sh.fill.fore_color.rgb))
        except:
            pass
    assert {"2E7D32", "E69138", "C0504D"} & fills
