from pptx import Presentation
from pptx.util import Emu
from status_deck.theme import Theme
from status_deck.slides import kpi_tiles

def _s():
    p = Presentation()
    p.slide_width = Emu(12192000)
    p.slide_height = Emu(6858000)
    return p.slides.add_slide(p.slide_layouts[6])

def test_kpi_tiles_render():
    s = _s()
    theme = Theme.anyreach()
    kpi_tiles.render(s, {"title": "At a glance", "tiles": [
        {"value": "8,315", "label": "in-scope calls", "sub": "26.9% of CS"},
        {"value": "65%", "label": "projected containment", "sub": ""}]}, theme)
    texts = [sh.text_frame.text for sh in s.shapes if sh.has_text_frame]
    assert "8,315" in texts and "projected containment" in texts
    # the big value uses the shared display size and the indigo data colour
    val_box = next(sh for sh in s.shapes if sh.has_text_frame and sh.text_frame.text == "8,315")
    run = val_box.text_frame.paragraphs[0].runs[0]
    assert run.font.size.pt == theme.size_display
    assert run.font.color.rgb == theme.color("indigo")
