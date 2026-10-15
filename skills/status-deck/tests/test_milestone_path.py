from pptx import Presentation
from pptx.util import Emu
from status_deck.theme import Theme
from status_deck.slides import milestone_path

def _s():
    p = Presentation()
    p.slide_width = Emu(12192000)
    p.slide_height = Emu(6858000)
    return p.slides.add_slide(p.slide_layouts[6])

def test_milestone_path():
    s = _s()
    milestone_path.render(s, {"title": "Milestones", "milestones": [
        {"date": "21 Sep", "label": "Phase 1 live", "done": False},
        {"date": "12 Oct", "label": "Phase 2 live", "done": False}]}, Theme.anyreach())
    texts = [sh.text_frame.text for sh in s.shapes if sh.has_text_frame]
    assert "21 Sep" in texts and "Phase 1 live" in texts
