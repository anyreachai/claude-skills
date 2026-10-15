from pptx import Presentation
from pptx.util import Emu
from pptx.enum.shapes import MSO_SHAPE, MSO_SHAPE_TYPE
from status_deck.theme import Theme
from status_deck.slides import gantt


def _s():
    p = Presentation()
    p.slide_width = Emu(12192000)
    p.slide_height = Emu(6858000)
    return p.slides.add_slide(p.slide_layouts[6])


def _content():
    return {
        "title": "Delivery plan",
        "eyebrow": "roadmap",
        "week_axis": {"labels": ["3 Aug", "10 Aug", "17 Aug", "24 Aug", "31 Aug"]},
        "here_week": 2,
        "phases": [
            {"name": "Phase 1", "note": "Core intents", "bars": [
                {"kind": "dev", "start_week": 0, "span": 3, "label": "Dev"},
                {"kind": "test", "start_week": 3, "span": 2, "label": "Test"},
                {"kind": "live", "start_week": 4, "span": 1, "label": "Live"}]},
            {"name": "Phase 2", "note": "Returns", "bars": [
                {"kind": "dev", "start_week": 1, "span": 2, "label": "Dev"}]},
        ],
        "gates": [{"week": 4, "label": "Gate 1"}],
    }


def test_gantt_renders():
    s = _s()
    gantt.render(s, _content(), Theme.anyreach())
    texts = [sh.text_frame.text for sh in s.shapes if sh.has_text_frame]
    assert "Phase 1" in texts
    assert "Phase 2" in texts
    assert "3 Aug" in texts
    assert "Gate 1" in texts
    # a bar label and the phase note carry through
    assert any("Core intents" in x for x in texts)
    assert any("Dev" == x for x in texts)


def test_gantt_default_legend():
    s = _s()
    c = _content()
    c.pop("legend", None)
    gantt.render(s, c, Theme.anyreach())
    texts = [sh.text_frame.text for sh in s.shapes if sh.has_text_frame]
    assert any("Solid bars are development" in x for x in texts)


def test_gantt_custom_legend():
    s = _s()
    c = _content()
    c["legend"] = "Dev then test then live"
    gantt.render(s, c, Theme.anyreach())
    texts = [sh.text_frame.text for sh in s.shapes if sh.has_text_frame]
    assert any("Dev then test then live" in x for x in texts)


def test_gantt_here_marker_and_gate_diamond():
    s = _s()
    gantt.render(s, _content(), Theme.anyreach())
    texts = [sh.text_frame.text for sh in s.shapes if sh.has_text_frame]
    assert any("we are here" in x for x in texts)
    diamonds = [sh for sh in s.shapes
                if sh.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE and sh.auto_shape_type == MSO_SHAPE.DIAMOND]
    assert len(diamonds) == 1


def test_gantt_week_commencing_caption():
    s = _s()
    gantt.render(s, _content(), Theme.anyreach())
    texts = [sh.text_frame.text for sh in s.shapes if sh.has_text_frame]
    assert any("week commencing" in x for x in texts)
