from pptx import Presentation
from pptx.util import Emu
from status_deck.theme import Theme
from status_deck.slides import pie

def _s():
    p = Presentation()
    p.slide_width = Emu(12192000)
    p.slide_height = Emu(6858000)
    return p.slides.add_slide(p.slide_layouts[6])

def test_pie_and_progress():
    s = _s()
    pie.render(s, {"title": "Mix", "series": [{"label": "A", "value": 60}, {"label": "B", "value": 40}]}, Theme.anyreach())
    assert any(sh.has_chart for sh in s.shapes)
    s2 = _s()
    theme = Theme.anyreach()
    pie.render_progress(s2, {"title": "Containment", "percent": 65, "label": "projected"}, theme)
    texts = [sh.text_frame.text for sh in s2.shapes if sh.has_text_frame]
    assert "65%" in texts and any(sh.has_chart for sh in s2.shapes)
    # the big centre percent uses the shared display size, not an ad-hoc point size
    pct_box = next(sh for sh in s2.shapes if sh.has_text_frame and sh.text_frame.text == "65%")
    assert pct_box.text_frame.paragraphs[0].runs[0].font.size.pt == theme.size_display
