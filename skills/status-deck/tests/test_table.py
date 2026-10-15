from pptx import Presentation
from pptx.util import Emu
from status_deck.theme import Theme
from status_deck.slides import table

def _slide():
    p=Presentation(); p.slide_width=Emu(12192000); p.slide_height=Emu(6858000)
    return p.slides.add_slide(p.slide_layouts[6])

def test_table_renders_header_and_rows():
    s=_slide()
    table.render(s,{"title":"Pending inputs","columns":["Input","Owner","Status"],
                    "rows":[["A","Matt","Pending"],["B","Dave","Recorded"]]},Theme.anyreach())
    texts=[sh.text_frame.text for sh in s.shapes if sh.has_text_frame]
    assert "INPUT" in texts and "A" in texts and "Recorded" in texts
