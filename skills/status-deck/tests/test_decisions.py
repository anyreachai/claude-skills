from pptx import Presentation
from pptx.util import Emu
from status_deck.theme import Theme
from status_deck.slides import decisions

def _s():
    p=Presentation(); p.slide_width=Emu(12192000); p.slide_height=Emu(6858000)
    return p.slides.add_slide(p.slide_layouts[6])

def test_decisions_status_colour():
    s=_s()
    decisions.render(s,{"title":"Decisions","rows":[
        {"item":"Send APIs","decision":"Klaviyo mid Sept","recorded":"17 Aug","status":"Recorded"},
        {"item":"Sheets","decision":"Structure to finalise","recorded":"","status":"Pending"}]},Theme.anyreach())
    reds=[sh for sh in s.shapes if sh.has_text_frame and sh.text_frame.text=="Pending"]
    assert reds and str(reds[0].text_frame.paragraphs[0].runs[0].font.color.rgb)=="C0504D"
