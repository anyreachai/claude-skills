from pptx import Presentation
from pptx.util import Emu
from status_deck.theme import Theme
from status_deck import blocks


def _slide():
    p = Presentation()
    p.slide_width = Emu(12192000)
    p.slide_height = Emu(6858000)
    return p.slides.add_slide(p.slide_layouts[6])


def test_text_box_sets_text_and_font():
    s = _slide()
    t = Theme.anyreach()
    shp = blocks.text_box(s, left=100, top=100, width=500, height=200, text="Hello",
                          size=t.size_body, color=t.color("body"), bold=True, font=t.font_body)
    assert shp.text_frame.text == "Hello"
    assert shp.text_frame.paragraphs[0].runs[0].font.bold is True


def test_header_band_fill():
    s = _slide()
    t = Theme.anyreach()
    shp = blocks.header_band(s, t, top=200, height=210)
    assert str(shp.fill.fore_color.rgb) == "1F2547"


def test_rag_chip_color():
    s = _slide()
    t = Theme.anyreach()
    shp = blocks.rag_chip(s, t, "red", left=100, top=100)
    assert str(shp.fill.fore_color.rgb) == "C0504D"
