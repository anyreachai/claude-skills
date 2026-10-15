from pptx import Presentation
from pptx.util import Emu
from pptx.enum.shapes import MSO_SHAPE, MSO_SHAPE_TYPE
from status_deck.theme import Theme
from status_deck.slides import exec_summary


def _s():
    p = Presentation()
    p.slide_width = Emu(12192000)
    p.slide_height = Emu(6858000)
    return p.slides.add_slide(p.slide_layouts[6])


def _rounded_rects(slide):
    out = []
    for sh in slide.shapes:
        if sh.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE and sh.auto_shape_type == MSO_SHAPE.ROUNDED_RECTANGLE:
            out.append(sh)
    return out


def test_exec_summary():
    s = _s()
    exec_summary.render(s, {
        "headline": "Phase 1 dev complete, testing next",
        "status": "red",
        "points": ["Dev complete", "Testing Thursday", "Inputs outstanding"],
        "visual": "none",
    }, Theme.anyreach())
    texts = [sh.text_frame.text for sh in s.shapes if sh.has_text_frame]
    assert any("Phase 1 dev complete" in x for x in texts)
    # status pill is a rounded rectangle carrying the status text, not plain coloured text
    assert any(x.strip().upper().startswith("STATUS: RED") for x in texts)
    pill_texts = [sh.text_frame.text for sh in _rounded_rects(s) if sh.has_text_frame]
    assert any("STATUS: RED" in x.upper() for x in pill_texts)


def test_exec_summary_one_card_per_point():
    s = _s()
    points = ["Dev complete", "Testing Thursday", "Inputs outstanding"]
    exec_summary.render(s, {
        "headline": "Headline",
        "status": "amber",
        "points": points,
        "visual": "none",
    }, Theme.anyreach())
    # one rounded rect per point (the cards) plus the status pill
    assert len(_rounded_rects(s)) == len(points) + 1
    texts = [sh.text_frame.text for sh in s.shapes if sh.has_text_frame]
    joined = "\n".join(texts)
    for pt in points:
        assert pt in joined


def test_exec_summary_colon_point_has_bold_lead():
    s = _s()
    theme = Theme.anyreach()
    exec_summary.render(s, {
        "headline": "Headline",
        "status": "green",
        "points": ["Returns: built and tested this week"],
        "visual": "none",
    }, theme)
    # the lead line renders bold; the detail renders as a second paragraph
    card_frames = [sh.text_frame for sh in s.shapes if sh.has_text_frame and "Returns" in sh.text_frame.text]
    tf = next(t for t in card_frames)
    assert tf.text.startswith("Returns")
    assert "built and tested this week" in tf.text
    lead_run = tf.paragraphs[0].runs[0]
    assert lead_run.font.bold is True
    # the bold lead line is a coloured heading (indigo), the detail stays body colour
    assert lead_run.font.color.rgb == theme.color("indigo")
    assert tf.paragraphs[1].runs[0].font.color.rgb == theme.color("body")
    assert len(tf.paragraphs) >= 2


def test_exec_summary_caps_at_five_points():
    s = _s()
    points = [f"Point {i}" for i in range(8)]
    exec_summary.render(s, {
        "headline": "Headline",
        "status": "green",
        "points": points,
        "visual": "none",
    }, Theme.anyreach())
    texts = [sh.text_frame.text for sh in s.shapes if sh.has_text_frame]
    joined = "\n".join(texts)
    for i in range(5):
        assert f"Point {i}" in joined
    for i in range(5, 8):
        assert f"Point {i}" not in joined
    # five cards plus the pill
    assert len(_rounded_rects(s)) == 6


def test_exec_summary_with_visual_progress():
    s = _s()
    exec_summary.render(s, {
        "headline": "Headline",
        "status": "green",
        "points": ["Only point"],
        "visual": "progress",
        "visual_content": {
            "title": "Build progress",
            "items": [{"label": "Workstream A", "build": 80, "tested": 40}],
        },
    }, Theme.anyreach())
    texts = [sh.text_frame.text for sh in s.shapes if sh.has_text_frame]
    assert any("Workstream A" in x for x in texts)


def test_exec_summary_no_visual_when_none():
    s = _s()
    exec_summary.render(s, {
        "headline": "Headline",
        "status": "green",
        "points": ["Only point"],
        "visual": "none",
    }, Theme.anyreach())
    texts = [sh.text_frame.text for sh in s.shapes if sh.has_text_frame]
    assert any("Headline" in x for x in texts)
