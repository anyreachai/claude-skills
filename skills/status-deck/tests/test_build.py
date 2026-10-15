import pytest
from pptx import Presentation
from status_deck.build import build_deck
from status_deck.verify import lint_brief, has_blocking

def _brief():
    return {"meta":{"title":"Status","parties":["Anyreach"],"date":"19 Aug 2026","status":"amber",
                    "footer":"Anyreach","confidential":True},
            "executive_summary":{"headline":"On track","status":"amber","points":["a","b"],"visual":"none"},
            "slides":[{"type":"table","title":"Asks","columns":["Ask","Owner"],"rows":[["x","Matt"]]},
                      {"type":"progress_pie","title":"Containment","percent":65,"label":"projected"}],
            "appendix":[{"type":"table","title":"Detail","columns":["A"],"rows":[["1"]]}]}

def test_order_and_divider(tmp_path):
    out=build_deck(_brief(), str(tmp_path/"d.pptx"))
    prs=Presentation(out)
    # exec summary (1) + 2 body + divider + 1 appendix = 5 slides
    assert len(prs.slides._sldIdLst)==5
    first=[sh.text_frame.text for sh in prs.slides[0].shapes if sh.has_text_frame]
    assert any("On track" in t for t in first)

def test_blocking_lint_raises(tmp_path):
    b=_brief(); b["meta"]["title"]="bad—dash"
    with pytest.raises(ValueError):
        build_deck(b, str(tmp_path/"d.pptx"))
