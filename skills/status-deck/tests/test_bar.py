from pptx import Presentation
from pptx.util import Emu
from status_deck.theme import Theme
from status_deck.slides import bar

def _s():
    p = Presentation()
    p.slide_width = Emu(12192000)
    p.slide_height = Emu(6858000)
    return p.slides.add_slide(p.slide_layouts[6])

def test_bar_adds_chart():
    s = _s()
    bar.render(s, {"title": "Volume by phase", "categories": ["P1", "P2", "P3"],
                   "series": [{"name": "Calls", "values": [8315, 10644, 11928]}], "mode": "simple"}, Theme.anyreach())
    assert any(sh.has_chart for sh in s.shapes)


def test_bar_series_use_data_palette_in_order():
    s = _s()
    theme = Theme.anyreach()
    bar.render(s, {"title": "Volume by phase", "categories": ["P1", "P2"],
                   "series": [{"name": "A", "values": [1, 2]}, {"name": "B", "values": [3, 4]},
                              {"name": "C", "values": [5, 6]}], "mode": "stacked"}, theme)
    ch = next(sh.chart for sh in s.shapes if sh.has_chart)
    fills = [ser.format.fill.fore_color.rgb for ser in ch.series]
    assert fills == [theme.series_color(0), theme.series_color(1), theme.series_color(2)]
