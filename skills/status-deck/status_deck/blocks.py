from pptx.util import Emu, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor
from .theme import hex_to_rgb

_ALIGN = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}
_VAL = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE, "bottom": MSO_ANCHOR.BOTTOM}


def filled_rect(slide, left, top, width, height, fill):
    sp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(left), Emu(top), Emu(width), Emu(height))
    sp.fill.solid()
    sp.fill.fore_color.rgb = fill if isinstance(fill, RGBColor) else hex_to_rgb(fill)
    sp.line.fill.background()
    sp.shadow.inherit = False
    return sp


def text_box(slide, *, left, top, width, height, text, size, color, bold=False, align="left", valign="top", font="Arial"):
    tb = slide.shapes.add_textbox(Emu(left), Emu(top), Emu(width), Emu(height))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = _VAL[valign]
    tf.margin_left = Emu(0)
    tf.margin_right = Emu(0)
    tf.margin_top = Emu(0)
    tf.margin_bottom = Emu(0)
    p = tf.paragraphs[0]
    p.alignment = _ALIGN[align]
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.name = font
    r.font.color.rgb = color if isinstance(color, RGBColor) else hex_to_rgb(color)
    return tb


def header_band(slide, theme, top, height):
    return filled_rect(slide, 325650, top, 11540700, height, theme.band_fill)


def band_label(slide, theme, text, left, top, width, height):
    return text_box(slide, left=left, top=top, width=width, height=height, text=text,
                    size=theme.size_table_header, color=hex_to_rgb(theme.band_text), bold=True,
                    valign="middle", font=theme.font_head)


def rag_chip(slide, theme, status, left, top, width=560000, height=200000):
    sp = filled_rect(slide, left, top, width, height, {"red": theme.rag_red, "amber": theme.rag_amber, "green": theme.rag_green}[status])
    return sp


def footnote(slide, theme, text, top):
    return text_box(slide, left=365760, top=top, width=11460480, height=260000, text=text,
                    size=theme.size_footer, color=theme.color("muted"), font=theme.font_body)
