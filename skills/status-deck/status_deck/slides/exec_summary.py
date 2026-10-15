from pptx.util import Emu, Pt
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

from .. import blocks
from . import gantt, progress_bars, milestone_path

LEFT = 342909
RIGHT = 11849400
CONTENT_W = RIGHT - LEFT

_VISUAL_GENERATORS = {
    "gantt": gantt,
    "progress": progress_bars,
    "milestone": milestone_path,
}

# pill
_PILL_TOP = 300000
_PILL_W = 2260000
_PILL_H = 380000
# headline
_HEAD_TOP = 840000
_HEAD_H = 1120000
# cards band (row of content cards sits in the upper-middle, below the headline)
_CARDS_TOP = 2200000
_CARD_GAP = 200000
_CARD_PAD = 180000
_MAX_POINTS = 5
_PER_ROW = 4
# card height is fit to wrapped content: estimate lines from text width, not a fixed box
_PT_EMU = 12700                     # 1 point in EMU
_CHAR_W_FACTOR = 0.56               # avg Arial glyph width as a fraction of the point size (conservative)
_LINE_H_FACTOR = 1.3                # line height as a fraction of the point size
_BODY_GAP = 4 * _PT_EMU             # matches the Pt(4) space_before on the detail paragraph
_MIN_CARD_H = 600000                # a one-line card still reads as a card


def _wrapped_lines(text, size_pt, inner_w):
    if not text:
        return 0
    char_w = _CHAR_W_FACTOR * size_pt * _PT_EMU
    chars_per_line = max(1, int(inner_w / char_w))
    return max(1, -(-len(text) // chars_per_line))  # ceil division


def _card_height(theme, point, inner_w):
    line_h = _LINE_H_FACTOR * theme.size_body * _PT_EMU
    if ":" in point:
        lead, rest = point.split(":", 1)
        lead, rest = lead.strip(), rest.strip()
        h = _wrapped_lines(lead, theme.size_body, inner_w) * line_h
        if rest:
            h += _BODY_GAP + _wrapped_lines(rest, theme.size_body, inner_w) * line_h
    else:
        h = _wrapped_lines(point, theme.size_body, inner_w) * line_h
    return max(_MIN_CARD_H, int(h + 2 * _CARD_PAD))


def _status_pill(slide, theme, status):
    sp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                Emu(LEFT), Emu(_PILL_TOP), Emu(_PILL_W), Emu(_PILL_H))
    sp.fill.solid()
    sp.fill.fore_color.rgb = theme.rag(status)
    sp.line.fill.background()
    sp.shadow.inherit = False
    tf = sp.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = Emu(0)
    tf.margin_right = Emu(0)
    tf.margin_top = Emu(0)
    tf.margin_bottom = Emu(0)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = f"STATUS: {status.upper()}"
    r.font.size = Pt(theme.size_section)
    r.font.bold = True
    r.font.name = theme.font_head
    r.font.color.rgb = theme.color("band_text")
    return sp


def _card(slide, theme, left, top, width, height, point):
    sp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                Emu(int(left)), Emu(int(top)), Emu(int(width)), Emu(int(height)))
    sp.fill.solid()
    sp.fill.fore_color.rgb = theme.color("zebra")
    sp.line.color.rgb = theme.color("border")
    sp.line.width = Emu(9525)
    sp.shadow.inherit = False

    tb = slide.shapes.add_textbox(Emu(int(left + _CARD_PAD)), Emu(int(top + _CARD_PAD)),
                                  Emu(int(width - 2 * _CARD_PAD)), Emu(int(height - 2 * _CARD_PAD)))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.TOP
    tf.margin_left = Emu(0)
    tf.margin_right = Emu(0)
    tf.margin_top = Emu(0)
    tf.margin_bottom = Emu(0)

    if ":" in point:
        lead, rest = point.split(":", 1)
        lead, rest = lead.strip(), rest.strip()
        p = tf.paragraphs[0]
        r = p.add_run()
        r.text = lead
        r.font.size = Pt(theme.size_body)
        r.font.bold = True
        r.font.name = theme.font_head
        r.font.color.rgb = theme.color("indigo")
        if rest:
            p2 = tf.add_paragraph()
            p2.space_before = Pt(4)
            r2 = p2.add_run()
            r2.text = rest
            r2.font.size = Pt(theme.size_body)
            r2.font.name = theme.font_body
            r2.font.color.rgb = theme.color("body")
    else:
        p = tf.paragraphs[0]
        r = p.add_run()
        r.text = point
        r.font.size = Pt(theme.size_body)
        r.font.name = theme.font_body
        r.font.color.rgb = theme.color("body")
    return sp


def _layout_cards(slide, theme, points):
    n = len(points)
    if n == 0:
        return
    rows = [points[:_PER_ROW]] if n <= _PER_ROW else [points[:_PER_ROW], points[_PER_ROW:]]
    top = _CARDS_TOP
    for row in rows:
        cols = len(row)
        card_w = (CONTENT_W - _CARD_GAP * (cols - 1)) / cols
        inner_w = card_w - 2 * _CARD_PAD
        # even row: every card takes the height of the tallest card in the row
        row_h = max(_card_height(theme, point, inner_w) for point in row)
        for ci, point in enumerate(row):
            left = LEFT + ci * (card_w + _CARD_GAP)
            _card(slide, theme, left, top, card_w, row_h, point)
        top += row_h + _CARD_GAP


def render(slide, content, theme):
    status = content.get("status", "green")
    _status_pill(slide, theme, status)

    blocks.text_box(slide, left=LEFT, top=_HEAD_TOP, width=CONTENT_W, height=_HEAD_H,
                    text=content["headline"], size=theme.size_title,
                    color=theme.color("ink"), bold=True, font=theme.font_head)

    points = content.get("points", [])[:_MAX_POINTS]
    _layout_cards(slide, theme, points)

    visual = content.get("visual", "none")
    visual_content = content.get("visual_content")
    if visual != "none" and visual_content:
        generator = _VISUAL_GENERATORS[visual]
        generator.render(slide, {**visual_content, "title": visual_content.get("title", "")}, theme)
