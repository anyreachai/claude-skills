from pptx.util import Emu
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from .. import blocks
from ..geometry import CONTENT_TOP, TITLE_TOP

LEFT = 342909

def _title(slide, theme, txt):
    blocks.text_box(slide, left=LEFT, top=TITLE_TOP, width=11500000, height=430000, text=txt,
                    size=theme.size_title, color=theme.color("ink"), bold=True, font=theme.font_head)

def render(slide, content, theme):
    _title(slide, theme, content["title"])
    d = CategoryChartData()
    d.categories = [s["label"] for s in content["series"]]
    d.add_series("", [s["value"] for s in content["series"]])
    gf = slide.shapes.add_chart(XL_CHART_TYPE.PIE, Emu(LEFT), Emu(CONTENT_TOP), Emu(6000000), Emu(4600000), d)
    ch = gf.chart
    pts = ch.plots[0].series[0].points
    for i, pt in enumerate(pts):
        pt.format.fill.solid()
        pt.format.fill.fore_color.rgb = theme.series_color(i)

def render_progress(slide, content, theme):
    _title(slide, theme, content["title"])
    pct = int(content["percent"])
    d = CategoryChartData()
    d.categories = ["done", "rest"]
    d.add_series("", [pct, 100 - pct])
    gf = slide.shapes.add_chart(XL_CHART_TYPE.DOUGHNUT, Emu(LEFT), Emu(CONTENT_TOP), Emu(4200000), Emu(4200000), d)
    ch = gf.chart
    ch.has_legend = False
    pts = ch.plots[0].series[0].points
    pts[0].format.fill.solid()
    pts[0].format.fill.fore_color.rgb = theme.color("indigo")
    pts[1].format.fill.solid()
    pts[1].format.fill.fore_color.rgb = theme.color("zebra")
    blocks.text_box(slide, left=LEFT, top=2800000, width=4200000, height=600000, text=f"{pct}%",
                    size=theme.size_display, color=theme.color("indigo"), bold=True, align="center", font=theme.font_head)
    if content.get("label"):
        blocks.text_box(slide, left=LEFT, top=3450000, width=4200000, height=350000, text=content["label"],
                        size=theme.size_subtext, color=theme.color("muted"), align="center", font=theme.font_body)
