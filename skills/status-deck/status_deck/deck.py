from pptx import Presentation
from pptx.util import Emu
from . import blocks
from .theme import Theme, hex_to_rgb
from .geometry import EMU_W, EMU_H

class Deck:
    def __init__(self, theme: Theme):
        self.theme=theme; self.prs=Presentation()
        self.prs.slide_width=Emu(EMU_W); self.prs.slide_height=Emu(EMU_H)
    def add_slide(self):
        return self.prs.slides.add_slide(self.prs.slide_layouts[6])
    def title_block(self, slide, title, subtitle=None, eyebrow=None):
        t=self.theme
        if eyebrow:
            blocks.text_box(slide,left=342909,top=150000,width=11500000,height=170000,text=eyebrow.upper(),
                            size=t.size_label,color=t.color("purple"),bold=True,font=t.font_head)
        blocks.text_box(slide,left=342909,top=330000,width=11500000,height=430000,text=title,
                        size=t.size_title,color=t.color("ink"),bold=True,font=t.font_head)
        if subtitle:
            blocks.text_box(slide,left=342909,top=760000,width=11500000,height=250000,text=subtitle,
                            size=t.size_subtitle,color=t.color("muted"),font=t.font_body)
        self.footer(slide)
    def footer(self, slide):
        t=self.theme
        line=("CONFIDENTIAL   " if t.confidential else "")+t.footer+(("   Prepared "+t.date) if t.date else "")
        blocks.text_box(slide,left=342909,top=6560000,width=11500000,height=180000,text=line,
                        size=t.size_confidential,color=t.color("muted"),font=t.font_body)
    def appendix_divider(self, title="Appendix"):
        s=self.add_slide(); t=self.theme
        blocks.header_band(s,t,top=3150000,height=560000)
        blocks.text_box(s,left=342909,top=3230000,width=11500000,height=400000,text=title,
                        size=t.size_title,color=hex_to_rgb(t.band_text),
                        bold=True,font=t.font_head,valign="middle")
        return s
    def save(self, path): self.prs.save(path); return path
